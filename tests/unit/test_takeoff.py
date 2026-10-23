"""Duct takeoff from vector drawings.

The reference figures in the comments come from sheets 6-7 of the benchmark
project (1181-18-2024-Р ОВ2); the tests themselves run on synthetic sheets so
they stay reproducible without proprietary files.
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from dc_extraction.takeoff import (  # noqa: E402
    centrelines,
    connect,
    detect_scale,
    leaders,
    merge_trunks,
    normalise_size,
    size_dimensions,
    snap_width,
    write_ru_csv,
    write_xlsx,
)
from dc_extraction.takeoff.callouts import (  # noqa: E402
    EQUIPMENT_RE,
    SIZE_RE,
    SYSTEM_RE,
    nearest_run,
)
from dc_extraction.takeoff.geometry import (  # noqa: E402
    PT_TO_MM_AT_1_1,
    DuctRun,
    drop_axis_lines,
    drop_hatch_pairs,
    page_segments,
)
from helpers.pdfgen import DuctSegment, make_duct_plan_pdf  # noqa: E402

MM_PER_PT = PT_TO_MM_AT_1_1 * 100

BLUE = (0.0, 0.0, 1.0)
RED = (1.0, 0.0, 0.0)


def plan(tmp_path: Path, segments, callouts=None) -> pymupdf.Page:
    pdf = tmp_path / "plan.pdf"
    make_duct_plan_pdf(pdf, segments, callouts)
    doc = pymupdf.open(pdf)
    return doc[0]


class TestMeasurement:
    def test_straight_duct_measures_its_true_length_and_width(self, tmp_path):
        page = plan(tmp_path, [DuctSegment((0, 0), (10, 0), 500, BLUE)])
        runs = centrelines(page, scale=100)
        total = sum(r.length_m for r in runs)
        assert total == pytest.approx(10.0, abs=0.3)
        assert max(runs, key=lambda r: r.length_m).width_mm == pytest.approx(500, abs=15)

    def test_coloured_stroke_length_is_not_the_duct_length(self, tmp_path):
        """The naive sum over coloured lines double counts every wall."""
        page = plan(tmp_path, [DuctSegment((0, 0), (10, 0), 500, BLUE)])
        strokes = sum(
            ((x1 - x0) ** 2 + (y1 - y0) ** 2) ** 0.5
            for x0, y0, x1, y1, colour in page_segments(page) if colour == BLUE
        )
        measured = sum(r.length_m for r in centrelines(page, scale=100))
        assert strokes * 25.4 / 72 * 100 / 1000 == pytest.approx(2 * measured, rel=0.05)

    def test_scale_is_detected_from_standard_duct_sizes(self, tmp_path):
        page = plan(tmp_path, [
            DuctSegment((0, 0), (8, 0), 600, BLUE),
            DuctSegment((0, 3), (8, 3), 250, BLUE),
        ])
        assert detect_scale(page) == 100

    def test_insulation_outlines_do_not_vote_for_the_scale(self, tmp_path):
        """Raw wall pairs of an insulated duct include its outer face (640 mm)
        and a wall-to-face band (620 mm); at 1:200 the latter reads as 1240,
        within tolerance of a standard 1250, and outvoted the true 1:100."""
        red = (1.0, 0.0, 0.0)
        page = plan(tmp_path, [
            DuctSegment((0, 0), (12, 0), 600, red, insulated=True),
            DuctSegment((8, 0), (8, 5), 400, red, insulated=True),
            DuctSegment((0, 9), (14, 9), 600, BLUE),
        ])
        assert detect_scale(page) == 100

    def test_insulated_duct_is_flagged_and_measured_at_duct_size(self, tmp_path):
        page = plan(tmp_path, [DuctSegment((0, 0), (12, 0), 600, RED, insulated=True)])
        runs = centrelines(page, scale=100)
        longest = max(runs, key=lambda r: r.length_m)
        assert longest.insulated
        # the outer face of the insulation is 640 mm; the duct is 600
        assert longest.width_mm == pytest.approx(600, abs=25)

    def test_plain_duct_is_not_flagged_as_insulated(self, tmp_path):
        page = plan(tmp_path, [DuctSegment((0, 0), (12, 0), 600, RED)])
        runs = centrelines(page, scale=100)
        assert not max(runs, key=lambda r: r.length_m).insulated

    @pytest.mark.parametrize("start, end", [
        ((0, 0), (12, 0)),
        # the same duct elsewhere on the sheet once surfaced a second run 20 mm
        # wider (inner wall of one side paired with the outer face of the other)
        ((0, 3), (12, 3)),
        ((0, 9), (12, 9)),
        ((8, 0), (8, 12)),
    ])
    def test_insulated_duct_is_counted_once(self, tmp_path, start, end):
        """Two concentric outlines are one duct, not two."""
        page = plan(tmp_path, [DuctSegment(start, end, 600, RED, insulated=True)])
        runs = centrelines(page, scale=100)
        assert sum(r.length_m for r in runs) == pytest.approx(12.0, abs=1.0)
        assert len(runs) == 1


class TestDrawingArtefacts:
    """Two things on a real sheet that pairing mistakes for ducts.

    Both are checked directly on constructed geometry: the synthetic plan
    generator does not reproduce either failure faithfully, and a test that
    passes with the fix reverted proves nothing.
    """

    def test_axis_line_of_a_round_duct_is_dropped(self):
        """A Ø200 carries a dash-dot centre line. Pairing takes the nearest
        partner on each side, so the axis wins over the far wall and the duct
        reads as two phantom 100 mm ducts - on sheets 6-7 that inflated the
        100 mm bucket to 411 m against 154 m taken off by hand."""
        wall_gap = 200 / MM_PER_PT          # Ø200 at 1:100
        segments = [
            (0.0, 0.0, 300.0, 0.0, 300.0),                        # wall
            (0.0, wall_gap / 2, 300.0, wall_gap / 2, 300.0),      # axis, midway
            (0.0, wall_gap, 300.0, wall_gap, 300.0),              # wall
        ]
        kept = drop_axis_lines(segments)
        assert len(kept) == 2
        assert all(abs(s[1] - wall_gap / 2) > 0.1 for s in kept), "the axis must go"

    def test_two_walls_without_an_axis_are_both_kept(self):
        wall_gap = 200 / MM_PER_PT
        segments = [
            (0.0, 0.0, 300.0, 0.0, 300.0),
            (0.0, wall_gap, 300.0, wall_gap, 300.0),
        ]
        assert len(drop_axis_lines(segments)) == 2

    def test_pair_crossed_by_a_hatch_family_is_rejected(self):
        """Inside an elbow the hatch fans out, so its strokes all differ in
        length and the equal-length seed misses them; two neighbours then pair
        into a duct exactly one hatch spacing wide. What still gives them away
        is the opposite family crossing at the lattice spacing."""
        spacing = 200 / MM_PER_PT
        line = (0.0, 0.0, 20 * spacing, 0.0)
        pair = {(0, 1): (spacing, line, 20 * spacing)}
        crossers = [
            (i * spacing, -10.0, i * spacing, 10.0, 20.0) for i in range(21)
        ]
        lattices = [(0.0, spacing), (math.pi / 2, spacing)]   # the cross-hatch
        assert drop_hatch_pairs(pair, lattices, crossers) == {}

    def test_a_real_duct_of_the_same_width_survives(self):
        """Same width, same direction, but only crossed here and there."""
        spacing = 200 / MM_PER_PT
        line = (0.0, 0.0, 20 * spacing, 0.0)
        pair = {(0, 1): (spacing, line, 20 * spacing)}
        crossers = [(3 * spacing, -10.0, 3 * spacing, 10.0, 20.0),
                    (11 * spacing, -10.0, 11 * spacing, 10.0, 20.0)]
        lattices = [(0.0, spacing), (math.pi / 2, spacing)]
        assert drop_hatch_pairs(pair, lattices, crossers) == pair


class TestNetwork:
    def test_elbow_joins_two_legs(self, tmp_path):
        """Arcs reach the page as sub-2pt chords and are dropped, so the corner
        has to be rebuilt from the two centrelines."""
        page = plan(tmp_path, [
            DuctSegment((0, 0), (8, 0), 600, BLUE),
            DuctSegment((8, 0), (8, 6), 600, BLUE),
        ])
        runs = centrelines(page, scale=100)
        trunks, _ = merge_trunks(runs)
        components = connect(trunks, scale=100)
        assert len(set(components)) == 1

    def test_crossing_ducts_stay_separate(self, tmp_path):
        """Two ducts crossing in plan are at different levels, not connected."""
        page = plan(tmp_path, [
            DuctSegment((0, 4), (16, 4), 400, BLUE),
            DuctSegment((8, 0), (8, 9), 400, RED),
        ])
        runs = centrelines(page, scale=100)
        trunks, _ = merge_trunks(runs)
        components = connect(trunks, scale=100)
        assert len(set(components)) >= 2

    def test_different_colours_never_merge(self, tmp_path):
        page = plan(tmp_path, [
            DuctSegment((0, 0), (8, 0), 500, BLUE),
            DuctSegment((8, 0), (8, 5), 500, RED),
        ])
        runs = centrelines(page, scale=100)
        trunks, _ = merge_trunks(runs)
        components = connect(trunks, scale=100)
        assert len(set(components)) >= 2


class TestCallouts:
    def test_label_typing(self):
        assert SIZE_RE.match("600x350")
        assert SIZE_RE.match("Ø200")
        assert SYSTEM_RE.match("ПД4.1")
        assert SYSTEM_RE.match("ДВ1а")
        assert not SYSTEM_RE.match("АМР")
        assert EQUIPMENT_RE.match("КЛАПЪ-(60)-НО-1100x600-BМ220")

    def test_size_spellings_normalise_to_one_form(self):
        assert normalise_size("ø200") == normalise_size("Ф200") == "Ø200"
        assert normalise_size("400х200") == "400x200"  # Cyrillic х
        assert size_dimensions("Ø200") == (200, 200)
        assert size_dimensions("600x350") == (600, 350)

    def test_stacked_callout_lines_share_one_leader(self, tmp_path):
        page = plan(
            tmp_path,
            [DuctSegment((0, 0), (10, 0), 600, BLUE)],
            callouts=[((500.0, 300.0), ["В4", "600x350"], (5.0, 0.0))],
        )
        found = leaders(page)
        kinds = {(kind, text) for kind, text, _, _ in found}
        assert ("system", "В4") in kinds
        assert ("size", "600x350") in kinds
        targets = {(round(x), round(y)) for _, _, x, y in found}
        assert len(targets) == 1, "both lines must resolve to the same duct"

    def test_grille_face_size_is_not_read_as_a_duct_section(self, tmp_path):
        """"АМР 400x200" names a grille's own face, not the duct it sits on -
        and a grille's face commonly matches the spigot it's mounted on, so
        the downstream width-match sanity check cannot catch this either. On
        sheet 6, 57 of 151 grille sizes slipped through undetected this way."""
        page = plan(
            tmp_path,
            [DuctSegment((0, 0), (10, 0), 400, BLUE)],
            callouts=[((500.0, 300.0), ["АМР 400x200"], (5.0, 0.0))],
        )
        found = leaders(page)
        assert not any(kind == "size" for kind, *_ in found)

    def test_a_genuine_duct_size_survives_without_a_grille_label(self, tmp_path):
        page = plan(
            tmp_path,
            [DuctSegment((0, 0), (10, 0), 400, BLUE)],
            callouts=[((500.0, 300.0), ["400x200"], (5.0, 0.0))],
        )
        found = leaders(page)
        assert any(kind == "size" and text == "400x200" for kind, text, _, _ in found)

    def test_leader_does_not_cross_over_to_a_wider_duct_of_another_family(self):
        """Once a leader's endpoint lands inside a band, distance-to-wall
        clamps to zero - a wide duct passing nearby then ties a slim one and,
        with no colour filter, wins by coming first in the list. On sheet 6
        this let a "ПД5.2" callout for a small orange damper get credited to
        an unrelated 450 mm blue duct routed past it, inflating ПД5.2 to 12.2 m
        against a 9.4 m ceiling in the specification."""
        wide_blue = DuctRun(colour="blue", width_mm=450.0, insulated=False,
                            p0=(0.0, 0.0), p1=(20.0, 0.0), length_m=20.0)
        slim_orange = DuctRun(colour="orange", width_mm=200.0, insulated=False,
                              p0=(10.0, -5.0), p1=(10.0, 5.0), length_m=10.0)
        runs = [wide_blue, slim_orange]
        # the leader lands exactly where the two ducts cross in plan - zero
        # raw distance and a zero clamped score against both
        index, _ = nearest_run(runs, 10.0, 0.0, scale=100)
        assert index == 0, "unfiltered, the earlier (wider) run wins the tie"
        index, _ = nearest_run(runs, 10.0, 0.0, scale=100, colours={"orange"})
        assert index == 1, "restricted to its own family, the real target wins"


class TestRuCsv:
    """Excel with Russian regional settings splits fields on ';' and expects
    ',' as the decimal mark; a plain csv.writer with '.' floats silently
    imports every number as text."""

    def test_floats_use_a_comma_decimal(self, tmp_path):
        out = write_ru_csv(tmp_path / "r.csv", ["Система", "Длина, м"],
                           [["В1", 154.4], ["В2", 0.0]])
        lines = out.read_text(encoding="utf-8-sig").splitlines()
        assert "154,4" in lines[1]
        assert "." not in "\n".join(lines[1:])

    def test_fields_are_semicolon_delimited(self, tmp_path):
        out = write_ru_csv(tmp_path / "r.csv", ["A", "B"], [[1.0, 2.0]])
        first_row = out.read_text(encoding="utf-8-sig").splitlines()[0]
        assert first_row == "A;B"


class TestXlsx:
    """The estimator's deliverable is a workbook: numbers must be numbers (so
    SUM works) and the totals row must be a live formula, not a pasted value."""

    def _open(self, path):
        import openpyxl
        return openpyxl.load_workbook(path)

    def test_numbers_stay_numeric(self, tmp_path):
        out = write_xlsx(tmp_path / "t.xlsx", "Ведомость", ["Система", "Длина, м"],
                         [["В1", 154.4], ["В2", 0.0]])
        ws = self._open(out)["Ведомость"]
        assert ws["A1"].value == "Система"
        assert ws["B2"].value == 154.4 and isinstance(ws["B2"].value, float)

    def test_totals_row_is_a_sum_formula(self, tmp_path):
        out = write_xlsx(tmp_path / "t.xlsx", "Ведомость", ["Система", "Длина, м"],
                         [["В1", 1.5], ["В2", 2.5]], total_columns=["Длина, м"])
        ws = self._open(out)["Ведомость"]
        assert ws["A4"].value == "Итого"
        assert ws["B4"].value == "=SUM(B2:B3)"

    def test_booleans_become_words_and_header_is_frozen(self, tmp_path):
        out = write_xlsx(tmp_path / "t.xlsx", "S", ["Изоляция"], [[True], [False]])
        ws = self._open(out)["S"]
        assert [ws["A2"].value, ws["A3"].value] == ["Да", "Нет"]
        assert ws.freeze_panes == "A2"

    def test_creates_missing_directories(self, tmp_path):
        out = write_xlsx(tmp_path / "a" / "b" / "t.xlsx", "S", ["A"], [[1]])
        assert out.exists()


class TestWidthSnapping:
    def test_snaps_to_the_standard_ladder(self):
        assert snap_width(649.0) == 650
        assert snap_width(1101.0) == 1100

    def test_rejects_a_width_that_is_not_a_duct(self):
        assert snap_width(1750.0) is None
