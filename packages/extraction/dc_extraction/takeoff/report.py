"""CSV export for Russian-locale spreadsheet software.

Windows' Russian regional settings use `;` as the list separator and `,` as
the decimal separator. A file that mixes the two - `;`-delimited fields with
`.`-decimal numbers, which is what a naive `csv.writer` produces - opens in
Excel with every numeric column silently imported as text: no right
alignment, no `SUM()`, and every review of it looks broken at a glance even
though the underlying figures are correct.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any


def _cell(value: Any, decimals: int) -> str:
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, float):
        return f"{value:.{decimals}f}".replace(".", ",")
    return str(value)


def write_ru_csv(path: str | Path, header: list[str], rows: list[list[Any]],
                 decimals: int = 1) -> Path:
    """Write a `;`-delimited CSV with `,` decimals, ready for RU-locale Excel."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(header)
        for row in rows:
            writer.writerow([_cell(v, decimals) for v in row])
    return path


_DECIMAL_FORMAT = "0.0"
_MAX_COLUMN_WIDTH = 60


def write_xlsx(path: str | Path, title: str, header: list[str], rows: list[list[Any]],
               total_columns: list[str] | None = None) -> Path:
    """Write a one-sheet workbook with real numeric cells.

    Unlike the CSV, no locale guessing is involved: floats are stored as
    numbers, and `total_columns` get a trailing `=SUM()` row so the estimator
    can edit a figure and see the total follow.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]  # Excel's sheet-name limit

    sheet.append(header)
    for row in rows:
        sheet.append([("Да" if v else "Нет") if isinstance(v, bool) else v for v in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    last = len(rows) + 1
    if total_columns and rows:
        total_row = last + 1
        sheet.cell(row=total_row, column=1, value="Итого").font = Font(bold=True)
        for name in total_columns:
            letter = get_column_letter(header.index(name) + 1)
            cell = sheet[f"{letter}{total_row}"]
            cell.value = f"=SUM({letter}2:{letter}{last})"
            cell.font = Font(bold=True)

    for column in sheet.columns:
        for cell in column:
            if isinstance(cell.value, float):
                cell.number_format = _DECIMAL_FORMAT
        widest = max(len(str(c.value)) for c in column if c.value is not None)
        sheet.column_dimensions[column[0].column_letter].width = min(widest + 2,
                                                                    _MAX_COLUMN_WIDTH)
    workbook.save(path)
    return path
